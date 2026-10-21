"""Builds templates/Utilisation-Mapping-Template.xlsx.

The Power Automate flow reads these tables by name, so the table names and
column headers here must match the guide exactly.

Run:  uv run --with openpyxl python scripts/build_template.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "templates" / "Utilisation-Mapping-Template.xlsx"
DATE_FMT = "dd-mmm-yyyy"


def add_table(ws, name, headers, rows, widths):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if not rows:  # an Excel table needs at least one body row
        ws.append([None] * len(headers))
    last_col = chr(ord("A") + len(headers) - 1)
    table = Table(displayName=name, ref=f"A1:{last_col}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for col, width in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


wb = Workbook()

readme = wb.active
readme.title = "Read me"
readme.column_dimensions["A"].width = 110
lines = [
    ("Clockify → Teams utilisation report: mapping workbook", Font(bold=True, size=14)),
    ("", None),
    ("Keep this file in OneDrive for Business or SharePoint. The flow reads the tables below by name.", None),
    ("Do not rename the tables or their column headers. Replace the example rows with your own.", Font(bold=True)),
    ("", None),
    ("Resources  – one row per person. Email must match the person's Clockify login email.", None),
    ("             Team Lead must match a name in the TeamLeads table. Daily Hours = working hours per day (blank = 8).", None),
    ("TeamLeads  – one row per Team Lead. Channel Link: in Teams, click ⋯ next to the channel → Copy link.", None),
    ("Leave      – one row per leave period (From and To are both required; use the same date for one day).", None),
    ("             Put ALL in the Email column for a public holiday: no report is sent for that day.", None),
    ("History    – optional. Only used if you add the 'history log' step from the guide.", None),
    ("", None),
    ("Tip: never leave empty rows in the middle of a table. To add a person, type in the first empty row under the table;", None),
    ("Excel extends the table automatically.", None),
]
for text, font in lines:
    readme.append([text])
    if font:
        readme.cell(row=readme.max_row, column=1).font = font
readme.sheet_view.showGridLines = False

resources = wb.create_sheet("Resources")
add_table(
    resources,
    "Resources",
    ["Resource Name", "Email", "Team Lead", "Daily Hours"],
    [
        ["Ananya Iyer", "ananya.iyer@example.com", "Kavita Shah", 8],
        ["Rohan Desai", "rohan.desai@example.com", "Kavita Shah", 8],
        ["Meera Pillai", "meera.pillai@example.com", "Kavita Shah", 4],
        ["Tom Fletcher", "tom.fletcher@example.com", "Arjun Nair", 8],
        ["Sara Khan", "sara.khan@example.com", "Arjun Nair", 8],
        ["Luis Ortega", "luis.ortega@example.com", "Arjun Nair", 7.5],
    ],
    [24, 34, 22, 14],
)
hours_rule = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True)
hours_rule.error = "Daily Hours must be a number between 0 and 24."
resources.add_data_validation(hours_rule)
hours_rule.add("D2:D1000")

team_leads = wb.create_sheet("TeamLeads")
add_table(
    team_leads,
    "TeamLeads",
    ["Team Lead", "Channel Link"],
    [
        ["Kavita Shah", "PASTE the Teams channel link here (⋯ next to channel → Copy link)"],
        ["Arjun Nair", "PASTE the Teams channel link here (⋯ next to channel → Copy link)"],
    ],
    [22, 120],
)

leave = wb.create_sheet("Leave")
add_table(
    leave,
    "Leave",
    ["Email", "From", "To", "Note"],
    [
        ["ALL", date(2026, 12, 25), date(2026, 12, 25), "Public holiday (example)"],
        ["rohan.desai@example.com", date(2026, 11, 2), date(2026, 11, 6), "Annual leave (example)"],
    ],
    [34, 16, 16, 34],
)
for row in leave.iter_rows(min_row=2, min_col=2, max_col=3):
    for cell in row:
        cell.number_format = DATE_FMT
date_rule = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
date_rule.error = "Please enter a date, e.g. 20-Oct-2026."
leave.add_data_validation(date_rule)
date_rule.add("B2:C2000")
for col in "BC":
    for r in range(2, 2001):
        leave[f"{col}{r}"].number_format = DATE_FMT

history = wb.create_sheet("History")
add_table(
    history,
    "History",
    ["Date", "Team Lead", "Resource", "Projects", "Logged (h)", "Capacity (h)", "Utilisation %", "Status"],
    [],
    [14, 22, 24, 50, 12, 13, 14, 26],
)

for ws in (resources, team_leads, leave, history):
    for cell in ws[1]:
        cell.alignment = Alignment(vertical="center")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
